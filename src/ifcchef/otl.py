"""Inlezen van de OTL-Excel."""

from pathlib import Path

import openpyxl

from ifcchef import config

# objecttype -> namen van toegestane properties (beide genormaliseerd met ``normalize``)
Otl = dict[str, frozenset[str]]


def normalize(value: object) -> str:
    return str(value).strip().lower()


def load_otl(path: Path) -> Otl:
    """Lees per objecttype welke properties volgens de OTL zijn toegestaan.

    Rij ``OTL_HEADER_ROW`` bevat de propertynamen; elke rij daaronder met een objecttype
    in kolom ``OTL_OBJECTTYPE_COLUMN`` geeft met ``x`` of ``?`` aan welke properties horen.
    Rijen zonder objecttype (categorieregels) worden overgeslagen.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if config.OTL_SHEET not in workbook.sheetnames:
            raise ValueError(f"Tabblad '{config.OTL_SHEET}' niet gevonden in {path}.")
        rows = workbook[config.OTL_SHEET].iter_rows(values_only=True)
        header: tuple = ()
        otl: Otl = {}
        for row_number, row in enumerate(rows, start=1):
            if row_number < config.OTL_HEADER_ROW:
                continue
            if row_number == config.OTL_HEADER_ROW:
                header = row
                continue
            objecttype = _cell(row, config.OTL_OBJECTTYPE_COLUMN - 1)
            if objecttype is None or not normalize(objecttype):
                continue
            allowed = {
                normalize(header[i])
                for i, value in enumerate(row)
                if i < len(header)
                and header[i] is not None
                and value is not None
                and normalize(value) in config.OTL_ALLOWED_MARKS
            }
            otl[normalize(objecttype)] = frozenset(allowed)
        return otl
    finally:
        workbook.close()


def _cell(row: tuple, index: int) -> object:
    return row[index] if index < len(row) else None
