"""Testdata: een kleine OTL-Excel en een klein IFC4-model met de ILS-pset.

OTL (zelfde layout als ROOT-OTL.xlsx: koppen in rij 2, objecttype in kolom C):
  betonwand       -> Objecttype, Materiaal
  metselwerkwand  -> Objecttype, Kleur (met '?')
  ruimte          -> Objecttype, Ruimtenaam
"""

from pathlib import Path

import ifcopenshell
import ifcopenshell.api
import openpyxl
import pytest

HEADERS = {7: "Objecttype", 8: "Materiaal", 9: "Kleur", 10: "Brandwerendheid", 11: "Ruimtenaam"}
ROWS = [
    {2: "Wanden"},  # categorieregel: geen objecttype in kolom C
    {3: "Betonwand", 7: "x", 8: "x"},
    {3: "Metselwerkwand", 7: "x", 9: "?"},
    {3: "Ruimte", 7: "x", 11: "x"},
]


@pytest.fixture
def otl_path(tmp_path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "OTL"
    workbook.create_sheet("Kenmerken")
    sheet.cell(1, 6, "Entiteit")
    sheet.cell(2, 3, "Objecttype")
    sheet.cell(2, 4, "Objecttype (generiek)")
    for col, name in HEADERS.items():
        sheet.cell(2, col, name)
    for row_number, row in enumerate(ROWS, start=3):
        for col, value in row.items():
            sheet.cell(row_number, col, value)
    path = tmp_path / "otl.xlsx"
    workbook.save(path)
    return path


def _ils(model, product, **properties):
    pset = ifcopenshell.api.run("pset.add_pset", model, product=product, name="ILS")
    ifcopenshell.api.run("pset.edit_pset", model, pset=pset, properties=properties)
    return pset


@pytest.fixture
def model() -> ifcopenshell.file:
    model = ifcopenshell.api.run("project.create_file", version="IFC4")
    ifcopenshell.api.run("root.create_entity", model, ifc_class="IfcProject", name="Test")

    def create(ifc_class, name):
        return ifcopenshell.api.run("root.create_entity", model, ifc_class=ifc_class, name=name)

    beton = create("IfcWall", "beton")
    _ils(model, beton, Objecttype="betonwand", Materiaal="C30/37", Kleur="grijs")

    metsel = create("IfcWall", "metsel")
    _ils(model, metsel, Objecttype="Metselwerkwand ", Kleur="rood", Materiaal="baksteen")

    # Type gedeeld door beton en metsel: toegestaan is de vereniging (Materiaal en Kleur).
    wall_type = create("IfcWallType", "wandtype")
    _ils(model, wall_type, Materiaal="x", Kleur="y", Brandwerendheid="60")
    ifcopenshell.api.run(
        "type.assign_type", model, related_objects=[beton, metsel], relating_type=wall_type
    )

    create("IfcWall", "zonder-ils")

    zonder_objecttype = create("IfcWall", "zonder-objecttype")
    _ils(model, zonder_objecttype, Materiaal="hout")

    onbekend = create("IfcWall", "onbekend")
    _ils(model, onbekend, Objecttype="glaswand", Kleur="blauw")

    ruimte = create("IfcSpace", "ruimte")
    _ils(model, ruimte, Objecttype="Ruimte", Ruimtenaam="Kantoor", Kleur="wit")

    opening = create("IfcOpeningElement", "opening")
    _ils(model, opening, Objecttype="betonwand", Kleur="n.v.t.")
    return model


@pytest.fixture
def ifc_path(tmp_path: Path, model: ifcopenshell.file) -> Path:
    path = tmp_path / "voorbeeld.ifc"
    model.write(str(path))
    return path
